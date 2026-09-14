from __future__ import annotations
import os
import json
import html
from typing import Any, Optional, List, Dict
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from app.config import settings

class DocumentExporter:
    """
    Exports Video Intelligence outputs (Summaries, Transcripts, Key Moments) into PDF, DOCX, and SRT formats.
    """
    def export_pdf(self, video_title: str, summary_data: dict, key_moments: list, output_filename: str) -> str:
        os.makedirs(settings.EXPORTS_DIR, exist_ok=True)
        file_path = os.path.join(settings.EXPORTS_DIR, output_filename)
        doc = SimpleDocTemplate(
            file_path,
            pagesize=letter,
            rightMargin=40,
            leftMargin=40,
            topMargin=40,
            bottomMargin=40
        )
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            'TitleStyle',
            parent=styles['Heading1'],
            fontSize=20,
            leading=24,
            textColor=colors.HexColor('#4F46E5'),
            spaceAfter=8
        )
        sub_style = ParagraphStyle(
            'SubTitleStyle',
            parent=styles['Normal'],
            fontSize=11,
            leading=15,
            textColor=colors.HexColor('#374151'),
            spaceAfter=12
        )
        h2_style = ParagraphStyle(
            'H2Style',
            parent=styles['Heading2'],
            fontSize=13,
            leading=17,
            textColor=colors.HexColor('#0891B2'),
            spaceBefore=10,
            spaceAfter=4
        )
        body_style = ParagraphStyle(
            'BodyStyle',
            parent=styles['BodyText'],
            fontSize=10,
            leading=14,
            textColor=colors.HexColor('#1F2937'),
            spaceAfter=4
        )

        elements = []
        safe_title = html.escape(str(video_title or "Video Summary Report"))
        elements.append(Paragraph(f"ClipMind AI — Video Intelligence Summary", title_style))
        elements.append(Paragraph(f"<b>Asset:</b> {safe_title}", sub_style))
        elements.append(Spacer(1, 8))

        # TL;DR Executive Summary
        tldr_text = html.escape(str(summary_data.get("tldr") or "Executive summary not available."))
        elements.append(Paragraph("Executive Summary (TL;DR)", h2_style))
        elements.append(Paragraph(tldr_text, body_style))
        elements.append(Spacer(1, 8))

        # Key Takeaways
        takeaways = summary_data.get("key_takeaways", [])
        if takeaways:
            elements.append(Paragraph("Key Takeaways", h2_style))
            for point in takeaways:
                safe_point = html.escape(str(point))
                elements.append(Paragraph(f"• {safe_point}", body_style))
            elements.append(Spacer(1, 8))

        # Structured Chapters
        sections = summary_data.get("sections", [])
        if sections:
            elements.append(Paragraph("Detailed Chapters", h2_style))
            for sec in sections:
                title_txt = html.escape(str(sec.get("title") or sec.get("heading") or "Chapter"))
                time_rng = html.escape(str(sec.get("timeRange") or sec.get("timestamp_range") or "00:00"))
                summary_txt = html.escape(str(sec.get("summary") or ""))
                elements.append(Paragraph(f"<b>{title_txt}</b> ({time_rng})<br/>{summary_txt}", body_style))
                for pt in (sec.get("bulletPoints") or sec.get("points") or []):
                    elements.append(Paragraph(f"• {html.escape(str(pt))}", body_style))
                elements.append(Spacer(1, 4))
            elements.append(Spacer(1, 8))

        # Key Moments & Timestamps
        if key_moments:
            elements.append(Paragraph("Key Moments & Timestamps", h2_style))
            for km in key_moments:
                ts = html.escape(str(km.get("timestamp_str") or km.get("timestamp") or "00:00"))
                km_title = html.escape(str(km.get("title") or "Key Segment"))
                km_imp = html.escape(str(km.get("importance") or km.get("importance_score") or "High"))
                km_desc = html.escape(str(km.get("summary") or km.get("description") or km.get("title") or ""))
                
                elements.append(Paragraph(f"<b>[{ts}] {km_title}</b> ({km_imp})<br/>{km_desc}", body_style))
                elements.append(Spacer(1, 4))

        doc.build(elements)
        return file_path

    def export_content(
        self,
        format_type: str,
        video: Any,
        transcript: Optional[Any] = None,
        summary: Optional[Any] = None,
        key_moments: Optional[list] = None
    ) -> str:
        """
        Unified dispatch method for exporting video intelligence in PDF, DOCX, SRT, VTT, or TXT formats.
        """
        os.makedirs(settings.EXPORTS_DIR, exist_ok=True)
        vid_id = str(getattr(video, "id", "video"))
        video_title = getattr(video, "title", "Video Intelligence Summary")

        sum_data = {}
        if summary:
            sum_data = {
                "tldr": getattr(summary, "tldr", "") or "",
                "key_takeaways": getattr(summary, "key_takeaways", []) or [],
                "sections": getattr(summary, "sections", []) or []
            }
        else:
            sum_data = {
                "tldr": f"Video intelligence summary for {video_title}.",
                "key_takeaways": [],
                "sections": []
            }

        segments = getattr(transcript, "segments", []) if transcript else []

        km_list = []
        if key_moments:
            for km in key_moments:
                if isinstance(km, dict):
                    km_list.append(km)
                else:
                    km_list.append({
                        "timestamp_str": getattr(km, "timestamp_str", None) or getattr(km, "timestamp", "00:00"),
                        "title": getattr(km, "title", "Key Moment"),
                        "importance": getattr(km, "importance_score", "High"),
                        "description": getattr(km, "description", "")
                    })

        fmt = (format_type or "txt").lower().strip()
        if fmt == "pdf":
            return self.export_pdf(video_title, sum_data, km_list, f"{vid_id}_summary.pdf")
        elif fmt == "docx":
            return self.export_docx(video_title, sum_data, f"{vid_id}_study_guide.docx")
        elif fmt == "srt":
            return self.export_srt(segments, f"{vid_id}_subtitles.srt")
        elif fmt == "vtt":
            return self.export_vtt(segments, f"{vid_id}_subtitles.vtt")
        elif fmt == "txt":
            return self.export_txt(video_title, segments, sum_data, f"{vid_id}_transcript.txt")
        else:
            return self.export_txt(video_title, segments, sum_data, f"{vid_id}_summary.txt")

    def export_srt(self, segments: list, output_filename: str) -> str:
        file_path = os.path.join(settings.EXPORTS_DIR, output_filename)
        lines = []
        for idx, seg in enumerate(segments):
            start = self._format_srt_time(seg.get("start", 0))
            end = self._format_srt_time(seg.get("end", 0))
            text = seg.get("text", "")
            lines.append(f"{idx+1}\n{start} --> {end}\n{text}\n")

        with open(file_path, "w", encoding="utf-8") as f:
            f.write("\n".join(lines))
        return file_path

    def export_vtt(self, segments: list, output_filename: str) -> str:
        file_path = os.path.join(settings.EXPORTS_DIR, output_filename)
        lines = ["WEBVTT\n"]
        for idx, seg in enumerate(segments):
            start = self._format_vtt_time(seg.get("start", 0))
            end = self._format_vtt_time(seg.get("end", 0))
            speaker = seg.get("speaker", "Speaker 1")
            text = seg.get("text", "")
            lines.append(f"{idx+1}\n{start} --> {end}\n<v {speaker}>{text}\n")

        with open(file_path, "w", encoding="utf-8") as f:
            f.write("\n".join(lines))
        return file_path

    def export_txt(self, video_title: str, segments: list, summary_data: dict, output_filename: str) -> str:
        file_path = os.path.join(settings.EXPORTS_DIR, output_filename)
        content = f"ClipMind AI — Video Intelligence Report\nTitle: {video_title}\n\n"
        content += f"=== EXECUTIVE SUMMARY ===\n{summary_data.get('tldr', '')}\n\n"
        content += "=== KEY TAKEAWAYS ===\n"
        for pt in summary_data.get("key_takeaways", []):
            content += f"- {pt}\n"
        content += "\n=== FULL TRANSCRIPT ===\n"
        for seg in segments:
            ts = seg.get("timestamp", "00:00")
            spk = seg.get("speaker", "Speaker 1")
            content += f"[{ts}] {spk}: {seg.get('text', '')}\n"

        with open(file_path, "w", encoding="utf-8") as f:
            f.write(content)
        return file_path

    def export_docx(self, video_title: str, summary_data: dict, output_filename: str) -> str:
        file_path = os.path.join(settings.EXPORTS_DIR, output_filename)
        try:
            import docx
            from docx.shared import Inches, Pt, RGBColor

            doc = docx.Document()
            doc.add_heading("ClipMind AI — Video Study Guide", level=0)
            p_sub = doc.add_paragraph()
            run_sub = p_sub.add_run(f"Lecture / Video Title: {video_title}")
            run_sub.bold = True

            # Executive Summary
            doc.add_heading("Executive Summary (TL;DR)", level=1)
            doc.add_paragraph(summary_data.get("tldr", "No summary available."))

            # Key Takeaways
            takeaways = summary_data.get("key_takeaways", [])
            if takeaways:
                doc.add_heading("Key Takeaways", level=1)
                for pt in takeaways:
                    doc.add_paragraph(pt, style="List Bullet")

            # Detailed Sections
            sections = summary_data.get("sections", [])
            if sections:
                doc.add_heading("Detailed Topic Breakdown", level=1)
                for sec in sections:
                    title_text = sec.get("title") or sec.get("heading") or "Section"
                    time_range = sec.get("timeRange") or sec.get("timestamp_range") or "00:00"
                    doc.add_heading(f"{title_text} ({time_range})", level=2)
                    doc.add_paragraph(sec.get("summary", ""))
                    for bp in sec.get("bulletPoints", []) or sec.get("points", []):
                        doc.add_paragraph(bp, style="List Bullet")

            doc.save(file_path)
            return file_path
        except Exception:
            content = f"# ClipMind AI — Study Guide\n\nTitle: {video_title}\n\n"
            content += f"## Executive Summary\n{summary_data.get('tldr', '')}\n\n"
            content += "## Key Takeaways\n"
            for pt in summary_data.get("key_takeaways", []):
                content += f"- {pt}\n"
            content += "\n## Detailed Topics\n"
            for sec in summary_data.get("sections", []):
                content += f"### {sec.get('title')} ({sec.get('timeRange')})\n"
                content += f"{sec.get('summary')}\n"
                for bp in sec.get("bulletPoints", []):
                    content += f"  * {bp}\n"
                content += "\n"

            with open(file_path, "w", encoding="utf-8") as f:
                f.write(content)
            return file_path

    def _format_srt_time(self, seconds: float) -> str:
        millis = int((seconds - int(seconds)) * 1000)
        m, s = divmod(int(seconds), 60)
        h, m = divmod(m, 60)
        return f"{h:02d}:{m:02d}:{s:02d},{millis:03d}"

    def _format_vtt_time(self, seconds: float) -> str:
        millis = int((seconds - int(seconds)) * 1000)
        m, s = divmod(int(seconds), 60)
        h, m = divmod(m, 60)
        return f"{h:02d}:{m:02d}:{s:02d}.{millis:03d}"

exporter = DocumentExporter()

