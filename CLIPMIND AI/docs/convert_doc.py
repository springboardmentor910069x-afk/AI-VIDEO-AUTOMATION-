import os
import re
import html
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from reportlab.lib.pagesizes import letter
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether, HRFlowable, Image as RLImage
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.pdfgen import canvas
from PIL import Image as PILImage

MD_PATH = r"d:/SPRING BOARD/AI-VIDEO-AUTOMATION-/CLIPMIND AI/docs/COMPLETE_PROJECT_DOCUMENTATION.md"
DOCX_PATH = r"d:/SPRING BOARD/AI-VIDEO-AUTOMATION-/CLIPMIND AI/docs/COMPLETE_PROJECT_DOCUMENTATION.docx"
PDF_PATH = r"d:/SPRING BOARD/AI-VIDEO-AUTOMATION-/CLIPMIND AI/docs/COMPLETE_PROJECT_DOCUMENTATION.pdf"
DOCS_DIR = os.path.dirname(MD_PATH)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def parse_markdown_blocks(md_content):
    lines = md_content.splitlines()
    blocks = []
    i = 0
    n = len(lines)
    
    while i < n:
        line = lines[i]
        stripped = line.strip()
        
        # Empty line
        if not stripped:
            i += 1
            continue
            
        # Code block
        if stripped.startswith("```"):
            lang = stripped[3:].strip()
            code_lines = []
            i += 1
            while i < n and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            if i < n:
                i += 1 # skip closing ```
            blocks.append({
                "type": "code",
                "lang": lang,
                "content": "\n".join(code_lines)
            })
            continue
            
        # Image block ![alt](path)
        m_img = re.match(r'^!\[(.*?)\]\((.*?)\)$', stripped)
        if m_img:
            blocks.append({
                "type": "image",
                "alt": m_img.group(1).strip(),
                "src": m_img.group(2).strip()
            })
            i += 1
            continue
            
        # Table
        if stripped.startswith("|") and stripped.endswith("|"):
            table_lines = []
            while i < n and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            # Parse table rows
            rows = []
            for t_line in table_lines:
                # Check if separator row like |:---|:---|
                cells = [c.strip() for c in t_line.split("|")[1:-1]]
                if all(re.match(r'^:?-+:?$', c) for c in cells if c):
                    continue
                rows.append(cells)
            if rows:
                blocks.append({
                    "type": "table",
                    "rows": rows
                })
            continue
            
        # Horizontal Rule
        if re.match(r'^-{3,}$', stripped) or re.match(r'^\*{3,}$', stripped):
            blocks.append({"type": "hr"})
            i += 1
            continue
            
        # Headings
        m_head = re.match(r'^(#{1,6})\s+(.*)$', stripped)
        if m_head:
            level = len(m_head.group(1))
            text = m_head.group(2).strip()
            blocks.append({
                "type": "heading",
                "level": level,
                "text": text
            })
            i += 1
            continue
            
        # Bullet list item
        m_bullet = re.match(r'^[-*]\s+(.*)$', stripped)
        if m_bullet:
            blocks.append({
                "type": "bullet",
                "text": m_bullet.group(1).strip()
            })
            i += 1
            continue
            
        # Numbered list item
        m_num = re.match(r'^(\d+)\.\s+(.*)$', stripped)
        if m_num:
            blocks.append({
                "type": "numbered",
                "num": m_num.group(1),
                "text": m_num.group(2).strip()
            })
            i += 1
            continue
            
        # Normal paragraph (accumulate contiguous lines)
        para_lines = [line]
        i += 1
        while i < n:
            next_line = lines[i]
            next_strip = next_line.strip()
            if (not next_strip or 
                next_strip.startswith("```") or 
                re.match(r'^!\[.*?\]\(.*?\)$', next_strip) or
                (next_strip.startswith("|") and next_strip.endswith("|")) or
                re.match(r'^-{3,}$', next_strip) or
                re.match(r'^(#{1,6})\s+', next_strip) or
                re.match(r'^[-*]\s+', next_strip) or
                re.match(r'^\d+\.\s+', next_strip)):
                break
            para_lines.append(next_line)
            i += 1
        blocks.append({
            "type": "paragraph",
            "text": " ".join(l.strip() for l in para_lines)
        })
        
    return blocks

# ----------------- DOCX GENERATION -----------------

def add_styled_run(para, text):
    """Parses basic inline markdown formatting: **bold**, *italic*, `code`"""
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith("**") and token.endswith("**") and len(token) >= 4:
            run = para.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("*") and token.endswith("*") and len(token) >= 2:
            run = para.add_run(token[1:-1])
            run.italic = True
        elif token.startswith("`") and token.endswith("`") and len(token) >= 2:
            run = para.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(185, 28, 28)
        else:
            para.add_run(token)

def generate_docx(blocks, output_path):
    doc = Document()
    
    # Page setup
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)
        
        # Header / Footer
        header = section.header
        hp = header.paragraphs[0]
        hp.text = "ClipMind AI — Video Summarization & Key Moments Detection Platform"
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hp.style.font.size = Pt(8.5)
        hp.style.font.color.rgb = RGBColor(120, 120, 120)
        
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.text = "ClipMind AI Project Documentation | Architecture & Milestone Implementation"
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.style.font.size = Pt(8.5)
        fp.style.font.color.rgb = RGBColor(140, 140, 140)

    for block in blocks:
        b_type = block["type"]
        
        if b_type == "heading":
            level = block["level"]
            text = block["text"]
            clean_text = text.replace("**", "").replace("`", "")
            
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            
            if level == 1:
                p.paragraph_format.space_before = Pt(18)
                p.paragraph_format.space_after = Pt(8)
                run = p.add_run(clean_text)
                run.bold = True
                run.font.size = Pt(18)
                run.font.color.rgb = RGBColor(30, 58, 138)
            elif level == 2:
                p.paragraph_format.space_before = Pt(14)
                p.paragraph_format.space_after = Pt(6)
                run = p.add_run(clean_text)
                run.bold = True
                run.font.size = Pt(14)
                run.font.color.rgb = RGBColor(3, 105, 161)
            elif level == 3:
                p.paragraph_format.space_before = Pt(10)
                p.paragraph_format.space_after = Pt(4)
                run = p.add_run(clean_text)
                run.bold = True
                run.font.size = Pt(12)
                run.font.color.rgb = RGBColor(30, 41, 59)
            else:
                p.paragraph_format.space_before = Pt(8)
                p.paragraph_format.space_after = Pt(2)
                run = p.add_run(clean_text)
                run.bold = True
                run.italic = True
                run.font.size = Pt(11)
                run.font.color.rgb = RGBColor(71, 85, 105)
                
        elif b_type == "paragraph":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.15
            add_styled_run(p, block["text"])
            
        elif b_type == "image":
            rel_src = block["src"]
            full_path = os.path.normpath(os.path.join(DOCS_DIR, rel_src))
            if os.path.exists(full_path):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(8)
                p.paragraph_format.space_after = Pt(4)
                run = p.add_run()
                run.add_picture(full_path, width=Inches(5.5))
                
        elif b_type == "bullet":
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.12
            add_styled_run(p, block["text"])
            
        elif b_type == "numbered":
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.12
            add_styled_run(p, block["text"])
            
        elif b_type == "hr":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            p_run = p.add_run("―" * 45)
            p_run.font.color.rgb = RGBColor(203, 213, 225)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
        elif b_type == "code":
            table = doc.add_table(rows=1, cols=1)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False
            cell = table.cell(0, 0)
            cell.width = Inches(6.8)
            set_cell_background(cell, "F8FAFC")
            set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
            
            cp = cell.paragraphs[0]
            cp.paragraph_format.space_before = Pt(0)
            cp.paragraph_format.space_after = Pt(0)
            c_run = cp.add_run(block["content"])
            c_run.font.name = "Consolas"
            c_run.font.size = Pt(8.5)
            c_run.font.color.rgb = RGBColor(30, 41, 59)
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
            
        elif b_type == "table":
            rows = block["rows"]
            if not rows:
                continue
            num_rows = len(rows)
            num_cols = max(len(r) for r in rows)
            
            t = doc.add_table(rows=num_rows, cols=num_cols)
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            t.autofit = True
            
            for r_idx, row_data in enumerate(rows):
                is_header = (r_idx == 0)
                for c_idx in range(num_cols):
                    cell = t.cell(r_idx, c_idx)
                    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                    set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
                    
                    val = row_data[c_idx] if c_idx < len(row_data) else ""
                    cp = cell.paragraphs[0]
                    cp.paragraph_format.space_before = Pt(0)
                    cp.paragraph_format.space_after = Pt(0)
                    cp.paragraph_format.line_spacing = 1.05
                    
                    if is_header:
                        set_cell_background(cell, "1E3A8A")
                        add_styled_run(cp, val)
                        for run in cp.runs:
                            run.bold = True
                            run.font.size = Pt(9.5)
                            run.font.color.rgb = RGBColor(255, 255, 255)
                    else:
                        if r_idx % 2 == 1:
                            set_cell_background(cell, "F8FAFC")
                        else:
                            set_cell_background(cell, "FFFFFF")
                        add_styled_run(cp, val)
                        for run in cp.runs:
                            run.font.size = Pt(9)
                            
            doc.add_paragraph().paragraph_format.space_after = Pt(6)

    doc.save(output_path)
    print(f"DOCX created successfully at: {output_path}")

# ----------------- PDF GENERATION -----------------

class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super(NumberedCanvas, self).__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_header_footer(num_pages)
            super(NumberedCanvas, self).showPage()
        super(NumberedCanvas, self).save()

    def draw_header_footer(self, page_count):
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#64748B"))
        
        # Header (pages > 1)
        if self._pageNumber > 1:
            self.drawString(36, 756, "ClipMind AI: Video Summarization & Key Moments Detection Platform")
            self.setStrokeColor(colors.HexColor("#CBD5E1"))
            self.setLineWidth(0.5)
            self.line(36, 750, 576, 750)
            
        # Footer
        footer_text = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(576, 25, footer_text)
        self.drawString(36, 25, "ClipMind AI Project Documentation & Implementation Guide")
        self.setStrokeColor(colors.HexColor("#CBD5E1"))
        self.setLineWidth(0.5)
        self.line(36, 35, 576, 35)
        
        self.restoreState()

def md_inline_to_reportlab(text):
    text = html.escape(text)
    text = re.sub(r'\*\*(.*?)\*\*', r'<b>\1</b>', text)
    text = re.sub(r'(?<!\*)\*(?!\*)(.*?)(?<!\*)\*(?!\*)', r'<i>\1</i>', text)
    text = re.sub(r'`(.*?)`', r'<font name="Courier" color="#B91C1C">\1</font>', text)
    text = text.replace(r'\|', '|').replace(r'\_', '_')
    return text

def generate_pdf(blocks, output_path):
    doc = SimpleDocTemplate(
        output_path,
        pagesize=letter,
        leftMargin=36,
        rightMargin=36,
        topMargin=54,
        bottomMargin=54
    )
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=18,
        leading=22,
        textColor=colors.HexColor('#1E3A8A'),
        spaceBefore=14,
        spaceAfter=6,
        keepWithNext=True
    )
    h2_style = ParagraphStyle(
        'DocH2',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=13,
        leading=16,
        textColor=colors.HexColor('#0369A1'),
        spaceBefore=12,
        spaceAfter=5,
        keepWithNext=True
    )
    h3_style = ParagraphStyle(
        'DocH3',
        parent=styles['Heading3'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=14,
        textColor=colors.HexColor('#1E293B'),
        spaceBefore=9,
        spaceAfter=4,
        keepWithNext=True
    )
    h4_style = ParagraphStyle(
        'DocH4',
        parent=styles['Heading4'],
        fontName='Helvetica-BoldOblique',
        fontSize=10,
        leading=13,
        textColor=colors.HexColor('#475569'),
        spaceBefore=6,
        spaceAfter=3,
        keepWithNext=True
    )
    body_style = ParagraphStyle(
        'DocBody',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=12.5,
        textColor=colors.HexColor('#1E293B'),
        spaceAfter=4
    )
    bullet_style = ParagraphStyle(
        'DocBullet',
        parent=body_style,
        leftIndent=15,
        firstLineIndent=-10,
        spaceAfter=2.5
    )
    code_style = ParagraphStyle(
        'DocCode',
        parent=styles['Code'],
        fontName='Courier',
        fontSize=7.5,
        leading=9.5,
        textColor=colors.HexColor('#0F172A'),
        backColor=colors.HexColor('#F8FAFC'),
        borderColor=colors.HexColor('#CBD5E1'),
        borderWidth=0.5,
        borderPadding=6,
        spaceBefore=4,
        spaceAfter=6
    )

    story = []

    for block in blocks:
        b_type = block["type"]
        
        if b_type == "heading":
            level = block["level"]
            formatted = md_inline_to_reportlab(block["text"])
            if level == 1:
                story.append(Paragraph(formatted, title_style))
            elif level == 2:
                story.append(Paragraph(formatted, h2_style))
            elif level == 3:
                story.append(Paragraph(formatted, h3_style))
            else:
                story.append(Paragraph(formatted, h4_style))
                
        elif b_type == "paragraph":
            formatted = md_inline_to_reportlab(block["text"])
            story.append(Paragraph(formatted, body_style))
            
        elif b_type == "image":
            rel_src = block["src"]
            full_path = os.path.normpath(os.path.join(DOCS_DIR, rel_src))
            if os.path.exists(full_path):
                try:
                    with PILImage.open(full_path) as im:
                        w, h = im.size
                    max_w = 480
                    aspect = h / w
                    render_w = min(max_w, w)
                    render_h = render_w * aspect
                    if render_h > 300:
                        render_h = 300
                        render_w = render_h / aspect
                    img_flowable = RLImage(full_path, width=render_w, height=render_h)
                    img_flowable.hAlign = 'CENTER'
                    story.append(Spacer(1, 4))
                    story.append(img_flowable)
                    story.append(Spacer(1, 4))
                except Exception as e:
                    print(f"Error embedding image {full_path}: {e}")
            
        elif b_type == "bullet":
            formatted = md_inline_to_reportlab(f"• {block['text']}")
            story.append(Paragraph(formatted, bullet_style))
            
        elif b_type == "numbered":
            formatted = md_inline_to_reportlab(f"{block['num']}. {block['text']}")
            story.append(Paragraph(formatted, bullet_style))
            
        elif b_type == "hr":
            story.append(Spacer(1, 4))
            story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor('#CBD5E1'), spaceAfter=6, spaceBefore=4))
            
        elif b_type == "code":
            safe_code = html.escape(block["content"]).replace("\n", "<br/>")
            story.append(Paragraph(safe_code, code_style))
            
        elif b_type == "table":
            rows = block["rows"]
            if not rows:
                continue
            col_count = max(len(r) for r in rows)
            available_width = 540
            col_width = available_width / col_count
            
            table_data = []
            for r_idx, row in enumerate(rows):
                row_cells = []
                is_header = (r_idx == 0)
                for c_idx in range(col_count):
                    val = row[c_idx] if c_idx < len(row) else ""
                    f_val = md_inline_to_reportlab(val)
                    if is_header:
                        cell_p = Paragraph(f"<b>{f_val}</b>", ParagraphStyle('TH', parent=body_style, fontSize=8, leading=10, textColor=colors.white))
                    else:
                        cell_p = Paragraph(f_val, ParagraphStyle('TD', parent=body_style, fontSize=7.5, leading=9.5))
                    row_cells.append(cell_p)
                table_data.append(row_cells)
                
            t = Table(table_data, colWidths=[col_width]*col_count)
            t.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E3A8A')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
                ('TOPPADDING', (0, 0), (-1, -1), 3),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
                ('LEFTPADDING', (0, 0), (-1, -1), 4),
                ('RIGHTPADDING', (0, 0), (-1, -1), 4),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')])
            ]))
            story.append(Spacer(1, 4))
            story.append(t)
            story.append(Spacer(1, 6))

    doc.build(story, canvasmaker=NumberedCanvas)
    print(f"PDF created successfully at: {output_path}")

if __name__ == "__main__":
    print(f"Reading markdown from: {MD_PATH}")
    with open(MD_PATH, "r", encoding="utf-8") as f:
        md_text = f.read()
    
    blocks = parse_markdown_blocks(md_text)
    print(f"Parsed {len(blocks)} semantic blocks.")
    
    generate_docx(blocks, DOCX_PATH)
    generate_pdf(blocks, PDF_PATH)
    print("All conversions completed successfully.")
